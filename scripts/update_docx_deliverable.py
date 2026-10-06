"""Produit une copie corrigée d'EdC-02 avec annexes et images."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import os
import tempfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "EdC-02-Livrable.docx"
OUTPUT = ROOT / "EdC-02-Livrable-corrige.docx"
DIAGRAMS = ROOT / "docs" / "diagrams"


def replace_paragraphs(document: Document) -> None:
    replacements = [
        (lambda text: "OUVRIR LES .MD" in text,
         "Le dépôt GitHub contient le code, les livrables et les diagrammes Mermaid affichés nativement par GitHub."),
        (lambda text: text.startswith("Une doccumentation complémentaire"),
         "Une documentation complémentaire traite de la sobriété numérique et de l'accessibilité."),
        (lambda text: text.startswith("Ce Bloc 2 reprend"),
         "L'EdC-01 fourni décrit les besoins, personas, rôles, critères, backlog et le choix Scrum. Il ne précise ni Jira, ni AWS, ni cadence de sprint, ni budget chiffré. Sa rubrique Gantt est vide; la section dépenses énumère des catégories sans montant. Les hypothèses complémentaires d'EdC-02 sont donc identifiées comme telles et non comme des reprises vérifiables. La mention de Sarah, absente de la source et sans rôle défini, est retirée."),
        (lambda text: text.startswith("Sprint de deux semaines choisi"),
         "Le cadrage retenu pour le MVP est de dix semaines, organisées en dix sprints hebdomadaires. La durée est une hypothèse de planification EdC-02, pas une consigne de l'EdC-01."),
        (lambda text: text.startswith("Équipe-projet reprise de l’EDC-01") or text.startswith("Équipe-projet reprise de l'EDC-01"),
         "Ressources du MVP : chef de projet, data scientist, développeur données, développeur API/interface, UX designer, QA/DevOps et appui RSSI/DPO. Les ETP, jours et TJM sont explicités dans le tableau budgétaire ci-dessous; les capteurs sont hors MVP."),
        (lambda text: text.startswith("WBS, backlog priorisé"),
         "Le planning définit les phases P0 à P3, les dix sprints, les ressources, le Gantt et les critères de recette dans les annexes ci-dessous."),
        (lambda text: text.startswith("Budget : MVP"),
         "Le budget MVP est recalculé à 84 000 € de ressources humaines à partir du planning. Les coûts cloud et licences du MVP local sont nuls; les capteurs sont exclus et à chiffrer pour une phase ultérieure."),
        (lambda text: text.startswith("Outils collaboratifs : Jira"),
         "GitHub héberge le code et le workflow Actions exécute pytest. Aucun tableau Jira ou GitHub Project avec tickets réels n'est fourni; la capture de preuve C8 reste à joindre après configuration."),
        (lambda text: text.startswith("Diagramme d’architecture (/docs/02") or text.startswith("Diagramme d'architecture (/docs/02"),
         "Diagramme d'architecture : source Mermaid dans docs/diagrams/architecture.md; export PNG en annexe."),
        (lambda text: text.startswith("Diagramme de composants (/docs/02") or text.startswith("Diagramme de composants (/docs/02"),
         "Diagramme de composants : source Mermaid dans docs/diagrams/composants.md; export PNG en annexe."),
        (lambda text: text.startswith("Collecte SYNOP Météo-France"),
         "Collecte SYNOP Météo-France (2015–2024); les journées sous-couvertes sont conservées avec une cible inconnue et exclues de l'entraînement."),
        (lambda text: text.startswith("29 181 observations"),
         "29 181 observations agrégées en 3 653 jours; 90 jours ont moins de huit rapports pluie et restent non étiquetés. Parmi les 3 563 jours étiquetés, 806 sont pluvieux (22,6 %)."),
        (lambda text: text.startswith("/model-inf"), "/model-info"),
        (lambda text: text.startswith("Rappel – part de positifis") or text.startswith("Rappel - part de positifis"),
         "Rappel — part des jours pluvieux correctement détectés, seuil fixé sur validation (0,2124) : 0,745."),
        (lambda text: text.startswith("ROC-AUC"),
         "ROC-AUC — capacité à distinguer positifs/négatifs : 0,7664."),
        (lambda text: text.startswith("PR-AUC"),
         "PR-AUC — compromis précision/rappel (classes déséquilibrées) : 0,5043."),
        (lambda text: text.startswith("Brier"),
         "Brier — calibration des probabilités : 0,1601 (baseline taux de base : 0,1882)."),
        (lambda text: text.startswith("F1"),
         "F1 — équilibre précision/rappel au seuil validé : 0,5506."),
        (lambda text: text.startswith("Le MVP suit les principes du RGESN"),
         "Le MVP utilise une régression logistique sur CPU, des conversions vectorisées, une collecte incrémentale et SQLite. CodeCarbon 3.3.1 estime cette exécution d'entraînement à 0,0027 gCO2e; c'est une estimation logicielle du processus local, pas une mesure au compteur. AWS n'est pas déployé et le PDF EdC-01 ne prescrit pas ce fournisseur. L'étude compare AWS, OVHcloud et Scaleway sans extrapoler de PUE global à un site précis."),
        (lambda text: text.startswith("Accessibilité :") or text.startswith("Accessibilité:"),
         "Accessibilité : la cible est RGAA 4.1.2 et WCAG 2.2 niveau AA. Le MVP n'a pas été audité et aucune conformité n'est revendiquée. Les contrastes, la navigation clavier, les lecteurs d'écran, les graphiques, le zoom et les navigateurs restent à tester manuellement."),
    ]
    matches = [0] * len(replacements)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        for index, (predicate, replacement) in enumerate(replacements):
            if predicate(text):
                paragraph.clear()
                paragraph.add_run(replacement)
                matches[index] += 1

    for index, count in enumerate(matches):
        if count != 1:
            raise RuntimeError(f"Remplacement {index} : {count} paragraphes correspondants.")

    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().lower() in {
            "chef de projet,", "ingénieur iot,", "2 data scientists,",
            "ux designer,", "2 développeurs.",
        }:
            paragraph._element.getparent().remove(paragraph._element)

    document.add_heading("Annexes — planning, budget et diagrammes", level=1)
    document.add_paragraph(
        "Plan relatif de dix semaines. Les dates calendaires seront fixées après validation du lancement."
    )
    add_figure(document, DIAGRAMS / "planning-gantt.png", "Gantt du MVP — phases P0 à P3.")

    document.add_heading("Phases et sprints", level=2)
    sprint_rows = [
        ("S1 / P0 / 1 sem.", "Cadrage, backlog, risques, dépôt et critères d'acceptation.", "Chef de projet, développeur API, RSSI/DPO"),
        ("S2 / P1 / 1 sem.", "Source SYNOP, schéma SQLite, tests de parsing.", "Data scientist, développeur données"),
        ("S3 / P1 / 1 sem.", "Collecte incrémentale, agrégats, contrôles des données.", "Data scientist, développeur données"),
        ("S4 / P2 / 1 sem.", "Analyse exploratoire, split temporel, baseline.", "Data scientist"),
        ("S5 / P2 / 1 sem.", "Features, entraînement comparatif, calibration.", "Data scientist, développeur données"),
        ("S6 / P2 / 1 sem.", "Validation, choix du seuil, métriques et artefacts.", "Data scientist, QA/DevOps"),
        ("S7 / P2 / 1 sem.", "API typée, prédictions unitaires et groupées.", "Développeur API/interface, data scientist"),
        ("S8 / P3 / 1 sem.", "Interface de démonstration et revue accessibilité.", "Développeur API/interface, UX designer"),
        ("S9 / P3 / 1 sem.", "Tests, CI, documentation, export et diagrammes.", "QA/DevOps, développeurs, chef de projet"),
        ("S10 / P3 / 1 sem.", "Recette, bilan, présentation et transfert.", "Équipe MVP, chef de projet"),
    ]
    add_table(document, ("Sprint / phase / durée", "Contenu", "Ressources"), sprint_rows)
    document.add_paragraph(
        "P0 = cadrage et socle; P1 = données et qualité; P2 = modèle et services; "
        "P3 = recette et transfert."
    )

    document.add_heading("Budget du MVP", level=2)
    budget_rows = [
        ("Chef de projet", "0,2", "10", "600 €", "6 000 €"),
        ("Data scientist", "0,8", "40", "600 €", "24 000 €"),
        ("Développeur données", "0,8", "40", "550 €", "22 000 €"),
        ("Développeur API/interface", "0,6", "30", "550 €", "16 500 €"),
        ("UX designer", "0,2", "10", "500 €", "5 000 €"),
        ("QA / DevOps", "0,3", "15", "500 €", "7 500 €"),
        ("Appui RSSI / DPO", "0,1", "5", "600 €", "3 000 €"),
        ("Total RH", "3,0", "150", "", "84 000 €"),
    ]
    add_table(document, ("Rôle", "ETP moyen", "Jours", "TJM HT", "Coût HT"), budget_rows)
    document.add_paragraph(
        "Hypothèses de calcul, pas dépenses constatées ni chiffres d'EdC-01. Cloud : 0 € "
        "(exécution locale); licences : 0 € (outils open source); capteurs : 0 € dans le MVP, "
        "achat et déploiement à chiffrer sur devis après S10. Les anciens montants de 147 k€, "
        "550 k€, 50 k€, 8 k€ et 5 k€ sont retirés faute de source dans le PDF fourni."
    )

    document.add_heading("Indicateurs qualité", level=2)
    quality_rows = [
        ("ROC-AUC", "≥ 0,75", "0,7664", "Atteint"),
        ("Brier", "Inférieur à la baseline taux de base", "0,1601 contre 0,1882", "Atteint"),
        ("PR-AUC", "Suivi; taux de base 0,2496", "0,5043", "Mesuré"),
        ("Précision des alertes", "> 90 %", "0,4366", "Non atteint"),
        ("Satisfaction utilisateurs", "≥ 80 % / 50 testeurs", "Non mesurée", "Non mesuré"),
        ("RMSE / MAE du cumul", "Prévision de quantité", "Non applicable au classifieur", "Hors périmètre"),
    ]
    add_table(document, ("Indicateur", "Objectif", "Résultat test", "Statut"), quality_rows)

    document.add_heading("Diagrammes exportés", level=2)
    add_figure(document, DIAGRAMS / "architecture.png", "Architecture du MVP livré.")
    add_figure(document, DIAGRAMS / "composants.png", "Composants logiciels et dépendances.")
    add_figure(document, DIAGRAMS / "deploiement.png", "Déploiement local actuel et cible non déployée.")
    document.add_paragraph(
        "Sources Mermaid : docs/diagrams/architecture.md, composants.md, deploiement.md "
        "et planning-gantt.md."
    )
    document.add_heading("Déclaration d'assistance", level=2)
    document.add_paragraph(
        "Une IA générative a été utilisée comme aide à la correction du code et à la rédaction "
        "et mise en forme de ce complément. Les résultats chiffrés proviennent d'exécutions "
        "locales et les sources de l'étude d'hébergement sont citées. L'auteur reste responsable "
        "du contenu et doit adapter cette mention aux consignes de son établissement."
    )


def add_table(document: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def add_figure(document: Document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(path), width=Inches(6.3))
    shape._inline.docPr.set("descr", caption)
    document.add_paragraph(caption, style="Caption")


def optimize_embedded_logo(path: Path) -> None:
    namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    with ZipFile(path, "r") as source:
        parts = {name: source.read(name) for name in source.namelist()}

    logo = Image.open(BytesIO(parts["word/media/image1.tif"]))
    logo.thumbnail((420, 420), Image.Resampling.LANCZOS)
    png_buffer = BytesIO()
    logo.save(png_buffer, format="PNG", optimize=True)
    parts.pop("word/media/image1.tif")
    parts["word/media/image1.png"] = png_buffer.getvalue()

    for name, content in list(parts.items()):
        if name.endswith(".rels"):
            parts[name] = content.replace(b"image1.tif", b"image1.png")
    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for element in list(content_types):
        if element.tag == f"{{{namespace}}}Default" and element.attrib.get("Extension", "").lower() in {"tif", "tiff"}:
            content_types.remove(element)
    if not any(
        element.tag == f"{{{namespace}}}Default"
        and element.attrib.get("Extension", "").lower() == "png"
        for element in content_types
    ):
        ET.SubElement(content_types, f"{{{namespace}}}Default", {
            "Extension": "png", "ContentType": "image/png"
        })
    parts["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )

    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with ZipFile(temporary_path, "w", ZIP_DEFLATED) as output:
            for name, content in parts.items():
                output.writestr(name, content)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    replace_paragraphs(document)
    document.save(OUTPUT)
    optimize_embedded_logo(OUTPUT)
    Document(OUTPUT)
    with ZipFile(OUTPUT) as package:
        if package.testzip() is not None:
            raise RuntimeError("Archive DOCX invalide après mise à jour.")
    print(f"Document corrigé créé : {OUTPUT} ({OUTPUT.stat().st_size:,} octets)")


if __name__ == "__main__":
    main()